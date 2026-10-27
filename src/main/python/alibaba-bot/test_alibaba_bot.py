import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest import mock
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook


SCRIPT_PATH = Path(__file__).with_name("alibaba-bot.py")
SPEC = importlib.util.spec_from_file_location("alibaba_bot", SCRIPT_PATH)
bot = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(bot)

EXPECTED_SEARCH_TEXTS = [
    "cordless chainsaw",
    "cordless electric chain saw",
    "gasoline generator",
    "petrol generator",
]
EXPECTED_COLUMNS = [
    "companyId", "url", "name", "city", "country", "verified",
    "manufacturer_type", "years", "staff", "factory_size(m\u00b2)", "worth(USD)",
    "rating", "reviews",
]
DENIED_URL = (
    "https://bixi-intl.alicdn.com/punish/punish:resource:template:ICBUSpace:default_447831.html"
    "?action=deny&origin=https%3A%2F%2Fwww.alibaba.com%3A443%2Fsearch%2Fpage"
)


def read_workbook_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()


class AlibabaBotTests(unittest.TestCase):
    def setUp(self):
        temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(temporary_directory.cleanup)
        self.output_dir = Path(temporary_directory.name) / "output"
        patches = [
            mock.patch.object(bot, "OUTPUT_DIR", str(self.output_dir)),
            mock.patch.object(
                bot, "PROFILE_DIR_FALLBACK", str(Path(temporary_directory.name) / "profile")
            ),
            mock.patch("builtins.print"),
        ]
        for patch in patches:
            patch.start()
            self.addCleanup(patch.stop)

    def test_single_configured_search_creates_one_workbook(self):
        phrase = "only configured search"
        data = [{"companyId": "single-id", "url": "https://supplier.example/", "name": phrase}]
        with (
            mock.patch.object(bot, "SEARCH_TEXTS", [phrase]),
            mock.patch.object(bot, "scrape", return_value=data) as scrape_mock,
        ):
            self.assertEqual(bot.main(), 0)

        scrape_mock.assert_called_once_with(phrase)
        self.assertEqual(len(list(self.output_dir.glob("*.xlsx"))), 1)
        rows = read_workbook_rows(bot.build_output_path(1, phrase))
        self.assertEqual(rows[0], tuple(EXPECTED_COLUMNS))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:3], ("single-id", "https://supplier.example/", phrase))

    def test_search_url_round_trips_plain_text_without_changing_filters(self):
        for phrase in [
            "baby play mat",
            'C++ & tools?color=red#mat / 100% "safe"',
            "m\u0103nu\u0219i \u5a74\u513f",
        ]:
            with self.subTest(phrase=phrase):
                url = urlsplit(bot.build_search_url(phrase))
                query = parse_qs(url.query)
                self.assertEqual(url.scheme, "https")
                self.assertEqual(url.netloc, "www.alibaba.com")
                self.assertEqual(url.path, "/search/page")
                self.assertEqual(url.fragment, "")
                self.assertEqual(query["SearchText"], [phrase])
                self.assertEqual(query["SearchScene"], ["suppliers"])
                self.assertNotIn("verifiedManufactory", query)
                self.assertEqual(set(query), {"spm", "SearchScene", "SearchText"})

    def test_output_names_handle_collisions_unicode_and_long_phrases(self):
        cases = [
            (1, "Baby & Mat", "alibaba_01_baby_mat_suppliers.xlsx"),
            (2, "Baby / Mat", "alibaba_02_baby_mat_suppliers.xlsx"),
            (3, "\u5a74\u513f <>:\"/\\|?*", "alibaba_03_search_suppliers.xlsx"),
            (4, "a" * 100, "alibaba_04_" + "a" * 80 + "_suppliers.xlsx"),
        ]
        paths = []
        for index, phrase, expected_name in cases:
            with self.subTest(phrase=phrase):
                path = Path(bot.build_output_path(index, phrase))
                self.assertEqual(path.parent, self.output_dir)
                self.assertEqual(path.name, expected_name)
                paths.append(path)
        self.assertEqual(len(set(paths)), len(cases))

    def test_batch_exports_four_separate_workbooks_in_order_with_per_file_dedup(self):
        completed_searches = []

        def scrape(phrase):
            index = len(completed_searches) + 1
            if completed_searches:
                previous_path = bot.build_output_path(index - 1, completed_searches[-1])
                self.assertTrue(Path(previous_path).is_file())
            completed_searches.append(phrase)
            return [
                {"companyId": "shared-id", "url": "https://supplier.example/", "name": phrase},
                {"companyId": "shared-id", "url": "https://supplier.example/", "name": "duplicate"},
            ]

        with (
            mock.patch.object(bot, "SEARCH_TEXTS", EXPECTED_SEARCH_TEXTS),
            mock.patch.object(bot, "scrape", side_effect=scrape) as scrape_mock,
        ):
            self.assertEqual(bot.main(), 0)

        self.assertEqual(scrape_mock.call_args_list, [mock.call(p) for p in EXPECTED_SEARCH_TEXTS])
        self.assertEqual(len(list(self.output_dir.glob("*.xlsx"))), 4)
        for index, phrase in enumerate(EXPECTED_SEARCH_TEXTS, start=1):
            with self.subTest(phrase=phrase):
                rows = read_workbook_rows(bot.build_output_path(index, phrase))
                self.assertEqual(rows[0], tuple(EXPECTED_COLUMNS))
                self.assertEqual(len(rows), 2)
                self.assertEqual(rows[1][:3], ("shared-id", "https://supplier.example/", phrase))

    def test_failed_second_search_does_not_prevent_third_export(self):
        phrases = EXPECTED_SEARCH_TEXTS[:3]
        with (
            mock.patch.object(bot, "SEARCH_TEXTS", phrases),
            mock.patch.object(bot, "scrape", side_effect=[[], RuntimeError("navigation failed"), []])
            as scrape_mock,
        ):
            self.assertEqual(bot.main(), 1)

        self.assertEqual(scrape_mock.call_args_list, [mock.call(p) for p in phrases])
        self.assertTrue(Path(bot.build_output_path(1, phrases[0])).is_file())
        self.assertFalse(Path(bot.build_output_path(2, phrases[1])).exists())
        self.assertTrue(Path(bot.build_output_path(3, phrases[2])).is_file())
        self.assertEqual(len(list(self.output_dir.glob("*.xlsx"))), 2)

    def test_empty_search_result_creates_headers_only_workbook(self):
        phrase = "empty result"
        with (
            mock.patch.object(bot, "SEARCH_TEXTS", [phrase]),
            mock.patch.object(bot, "scrape", return_value=[]),
        ):
            self.assertEqual(bot.main(), 0)

        rows = read_workbook_rows(bot.build_output_path(1, phrase))
        self.assertEqual(rows, [tuple(EXPECTED_COLUMNS)])

    def test_invalid_search_configuration_stops_before_setup_or_scraping(self):
        for phrases in [[], ["  "], ["valid phrase", ""], ["valid phrase", None], "one phrase"]:
            with self.subTest(phrases=phrases):
                with (
                    mock.patch.object(bot, "SEARCH_TEXTS", phrases),
                    mock.patch.object(bot.importlib, "import_module") as import_module,
                    mock.patch.object(bot, "scrape") as scrape_mock,
                ):
                    self.assertEqual(bot.main(), 1)
                import_module.assert_not_called()
                scrape_mock.assert_not_called()
                self.assertFalse(self.output_dir.exists())

    def test_missing_excel_dependency_stops_before_scraping(self):
        executable = r"C:\Program Files\Python 3.14\python.exe"
        with (
            mock.patch("builtins.print") as print_mock,
            mock.patch.object(bot.sys, "executable", executable),
            mock.patch.object(bot, "scrape") as scrape_mock,
            mock.patch.object(
                bot.importlib, "import_module", side_effect=ModuleNotFoundError("openpyxl")
            ) as import_module,
        ):
            self.assertEqual(bot.main(), 1)

        import_module.assert_called_once_with("openpyxl")
        requirements_path = Path(bot.__file__).with_name("requirements.txt")
        print_mock.assert_any_call(
            f'Install dependencies with: & "{executable}" -m pip install -r "{requirements_path}"'
        )
        scrape_mock.assert_not_called()
        self.assertFalse(self.output_dir.exists())

    def test_scrape_closes_browser_when_navigation_fails(self):
        driver = mock.Mock()
        phrase = "chainsaw + oil"
        with (
            mock.patch.object(bot, "make_driver", return_value=driver),
            mock.patch.object(bot, "robust_go", return_value=False) as robust_go,
            mock.patch.object(bot, "run_flow") as run_flow,
        ):
            with self.assertRaises(RuntimeError):
                bot.scrape(phrase)

        robust_go.assert_called_once_with(driver, bot.build_search_url(phrase))
        run_flow.assert_not_called()
        driver.quit.assert_called_once_with()

    def test_scrape_closes_browser_when_collection_fails(self):
        driver = mock.Mock()
        with (
            mock.patch.object(bot, "make_driver", return_value=driver),
            mock.patch.object(bot, "robust_go", return_value=True),
            mock.patch.object(bot, "run_flow", side_effect=RuntimeError("collection failed")),
        ):
            with self.assertRaisesRegex(RuntimeError, "collection failed"):
                bot.scrape("chainsaw")

        driver.quit.assert_called_once_with()

    def test_setup_mode_only_opens_profile_without_scraping_or_export_setup(self):
        driver = mock.Mock()
        with (
            mock.patch.object(bot.sys, "argv", [str(SCRIPT_PATH), "--setup-browser"]),
            mock.patch.object(bot, "HEADLESS", False),
            mock.patch.object(bot, "SEARCH_TEXTS", []),
            mock.patch.object(bot, "make_driver", return_value=driver) as make_driver,
            mock.patch("builtins.input", return_value="") as input_mock,
            mock.patch.object(bot, "scrape") as scrape_mock,
            mock.patch.object(bot, "robust_go") as navigate_mock,
            mock.patch.object(bot.importlib, "import_module") as import_module,
        ):
            self.assertEqual(bot.main(), 0)

        make_driver.assert_called_once_with(
            bot.PROFILE_DIR_FALLBACK, bot.PROFILE_NAME_FALLBACK, "SETUP"
        )
        input_mock.assert_called_once()
        scrape_mock.assert_not_called()
        navigate_mock.assert_not_called()
        import_module.assert_not_called()
        driver.get.assert_not_called()
        driver.execute_script.assert_not_called()
        driver.quit.assert_called_once_with()
        self.assertFalse(self.output_dir.exists())

    def test_setup_mode_closes_browser_when_input_is_unavailable(self):
        driver = mock.Mock()
        with (
            mock.patch.object(bot, "HEADLESS", False),
            mock.patch.object(bot, "make_driver", return_value=driver),
            mock.patch("builtins.input", side_effect=EOFError),
        ):
            self.assertEqual(bot.setup_browser(), 1)

        driver.get.assert_not_called()
        driver.execute_script.assert_not_called()
        driver.quit.assert_called_once_with()
        self.assertFalse(self.output_dir.exists())

    def test_headless_setup_stops_before_opening_a_browser(self):
        with (
            mock.patch.object(bot, "HEADLESS", True),
            mock.patch.object(bot, "make_driver") as make_driver,
            mock.patch("builtins.input") as input_mock,
        ):
            self.assertEqual(bot.setup_browser(), 1)

        make_driver.assert_not_called()
        input_mock.assert_not_called()

    def test_browser_vpn_prompt_precedes_navigation_and_collection(self):
        driver = mock.Mock()
        events = []
        rows = [{"companyId": "supplier", "url": "https://supplier.example/"}]

        def prepare_vpn(_prompt):
            events.append("vpn-ready")
            return ""

        def navigate(_driver, _url):
            events.append("navigate")
            return True

        def collect(_driver):
            events.append("collect")
            return rows

        with (
            mock.patch.object(bot.sys, "argv", [str(SCRIPT_PATH), "--browser-vpn"]),
            mock.patch.object(bot, "HEADLESS", False),
            mock.patch.object(bot, "make_driver", return_value=driver),
            mock.patch("builtins.input", side_effect=prepare_vpn),
            mock.patch.object(bot, "robust_go", side_effect=navigate) as navigate_mock,
            mock.patch.object(bot, "run_flow", side_effect=collect),
        ):
            self.assertEqual(bot.scrape("chainsaw"), rows)

        self.assertEqual(events, ["vpn-ready", "navigate", "collect"])
        navigate_mock.assert_called_once_with(driver, bot.build_search_url("chainsaw"))
        driver.quit.assert_called_once_with()

    def test_cancelled_browser_vpn_preparation_closes_without_navigation(self):
        for input_error in [EOFError, KeyboardInterrupt]:
            with self.subTest(input_error=input_error.__name__):
                driver = mock.Mock()
                expected_error = (
                    bot.AlibabaAccessBlocked if input_error is EOFError else KeyboardInterrupt
                )
                with (
                    mock.patch.object(bot.sys, "argv", [str(SCRIPT_PATH), "--browser-vpn"]),
                    mock.patch.object(bot, "HEADLESS", False),
                    mock.patch.object(bot, "make_driver", return_value=driver),
                    mock.patch("builtins.input", side_effect=input_error),
                    mock.patch.object(bot, "robust_go") as navigate_mock,
                    mock.patch.object(bot, "run_flow") as collect_mock,
                    self.assertRaises(expected_error),
                ):
                    bot.scrape("chainsaw")

                navigate_mock.assert_not_called()
                collect_mock.assert_not_called()
                driver.get.assert_not_called()
                driver.execute_script.assert_not_called()
                driver.quit.assert_called_once_with()

    def test_denied_navigation_stops_after_one_request_even_when_get_raises(self):
        for navigation_error in [None, bot.WebDriverException("page load timed out")]:
            with self.subTest(navigation_error=navigation_error):
                driver = mock.Mock(current_url=DENIED_URL)
                driver.get.side_effect = navigation_error
                with (
                    mock.patch.object(bot.time, "sleep"),
                    mock.patch("builtins.input") as input_mock,
                    self.assertRaisesRegex(bot.AlibabaAccessBlocked, "no CAPTCHA"),
                ):
                    bot.robust_go(driver, bot.build_search_url("chainsaw"))

                driver.get.assert_called_once_with(bot.build_search_url("chainsaw"))
                driver.execute_script.assert_not_called()
                driver.find_elements.assert_not_called()
                input_mock.assert_not_called()

    def test_visible_unusual_traffic_warning_is_denied_without_captcha_prompt(self):
        driver = mock.Mock(current_url=bot.build_search_url("chainsaw"))
        driver.find_elements.return_value = [mock.Mock(text=(
            "We have detected unusual traffic from your network,\n"
            "please try again later."
        ))]
        with (
            mock.patch("builtins.input") as input_mock,
            self.assertRaises(bot.AlibabaAccessBlocked),
        ):
            bot.check_access(driver)

        input_mock.assert_not_called()
        driver.get.assert_not_called()
        driver.execute_script.assert_not_called()

    def test_navigation_does_not_accept_alibaba_domain_in_another_hosts_query(self):
        for url in [
            "https://example.com/?origin=https%3A%2F%2Fwww.alibaba.com%2Fsearch%2Fpage",
            "https://www.alibaba.com.example.com/",
        ]:
            with self.subTest(url=url):
                driver = mock.Mock(current_url=url)
                driver.find_elements.return_value = []
                with mock.patch.object(bot.time, "sleep"):
                    self.assertFalse(bot.robust_go(driver, bot.build_search_url("chainsaw"), attempts=1))

        for url in ["https://alibaba.com/", "https://www.alibaba.com/search/page"]:
            with self.subTest(url=url):
                driver = mock.Mock(current_url=url)
                driver.find_elements.return_value = []
                with mock.patch.object(bot.time, "sleep"):
                    self.assertTrue(bot.robust_go(driver, url))
                driver.get.assert_called_once_with(url)
                driver.execute_script.assert_not_called()

    def test_manual_verification_can_complete_without_automatic_navigation(self):
        driver = mock.Mock(current_url="https://www.alibaba.com/punish?action=captcha")
        body = mock.Mock(text="Slide to verify")
        driver.find_elements.return_value = [body]

        def complete_verification(_prompt):
            driver.current_url = bot.build_search_url("chainsaw")
            body.text = "Supplier results"
            return ""

        with (
            mock.patch.object(bot, "HEADLESS", False),
            mock.patch("builtins.input", side_effect=complete_verification) as input_mock,
        ):
            bot.check_access(driver)

        input_mock.assert_called_once()
        driver.get.assert_not_called()
        driver.execute_script.assert_not_called()

    def test_unfinished_manual_verification_stops_without_reloading(self):
        driver = mock.Mock(current_url="https://www.alibaba.com/punish?action=captcha")
        driver.find_elements.return_value = [mock.Mock(text="Slide to verify")]
        with (
            mock.patch.object(bot, "HEADLESS", False),
            mock.patch("builtins.input", return_value="") as input_mock,
            self.assertRaisesRegex(bot.AlibabaAccessBlocked, "still displayed"),
        ):
            bot.check_access(driver)

        input_mock.assert_called_once()
        driver.get.assert_not_called()
        driver.execute_script.assert_not_called()

    def test_initial_denial_stops_before_login_prompt_or_scrolling(self):
        driver = mock.Mock(current_url=DENIED_URL)
        with (
            mock.patch.object(bot, "WAIT_FOR_LOGIN", True),
            mock.patch("builtins.input") as input_mock,
            mock.patch.object(bot, "scroll_to_bottom") as scroll_mock,
            self.assertRaises(bot.AlibabaAccessBlocked) as raised,
        ):
            bot.run_flow(driver)

        self.assertEqual(raised.exception.rows, [])
        input_mock.assert_not_called()
        scroll_mock.assert_not_called()
        driver.execute_script.assert_not_called()

    def test_card_wait_does_not_scroll_a_denied_page(self):
        driver = mock.Mock(current_url=DENIED_URL)
        with (
            mock.patch.object(bot, "scroll_nudge") as nudge_mock,
            self.assertRaises(bot.AlibabaAccessBlocked),
        ):
            bot.wait_for_cards(driver)

        nudge_mock.assert_not_called()
        driver.execute_script.assert_not_called()

    def test_pagination_denial_preserves_already_collected_rows(self):
        driver = mock.Mock(current_url=bot.build_search_url("chainsaw"))
        driver.find_elements.return_value = []
        card, next_button = mock.Mock(), mock.Mock()
        row = {"companyId": "first", "url": "https://supplier.example/", "name": "Supplier"}

        def deny_next_page():
            driver.current_url = DENIED_URL

        next_button.click.side_effect = deny_next_page
        with (
            mock.patch.object(bot, "WAIT_FOR_LOGIN", False),
            mock.patch.object(bot.time, "sleep"),
            mock.patch.object(bot, "scroll_to_bottom"),
            mock.patch.object(bot, "wait_for_cards", return_value=[card]) as wait_mock,
            mock.patch.object(bot, "get_page_label", return_value=1),
            mock.patch.object(bot, "extract_from_card", return_value=row),
            mock.patch.object(bot, "find_next", return_value=next_button) as find_next_mock,
            self.assertRaises(bot.AlibabaAccessBlocked) as raised,
        ):
            bot.run_flow(driver)

        self.assertEqual(raised.exception.rows, [row])
        self.assertEqual(wait_mock.call_count, 2)
        find_next_mock.assert_called_once_with(driver)
        next_button.click.assert_called_once_with()
        driver.get.assert_not_called()

    def test_first_card_key_does_not_swallow_denial_or_discard_collected_rows(self):
        driver = mock.Mock(current_url=bot.build_search_url("chainsaw"))
        driver.find_elements.return_value = []
        card = mock.Mock()
        row = {"companyId": "first", "url": "https://supplier.example/"}
        with (
            mock.patch.object(bot, "WAIT_FOR_LOGIN", False),
            mock.patch.object(bot, "scroll_to_bottom"),
            mock.patch.object(bot, "wait_for_cards", side_effect=[
                [card], bot.AlibabaAccessBlocked("Access denied while checking first card")
            ]),
            mock.patch.object(bot, "get_page_label", return_value=1),
            mock.patch.object(bot, "extract_from_card", return_value=row),
            mock.patch.object(bot, "find_next") as find_next_mock,
            self.assertRaises(bot.AlibabaAccessBlocked) as raised,
        ):
            bot.run_flow(driver)

        self.assertEqual(raised.exception.rows, [row])
        find_next_mock.assert_not_called()

    def test_initial_block_stops_remaining_searches_without_creating_empty_workbook(self):
        phrases = EXPECTED_SEARCH_TEXTS[:3]
        with (
            mock.patch.object(bot, "SEARCH_TEXTS", phrases),
            mock.patch.object(bot, "scrape", side_effect=bot.AlibabaAccessBlocked("Access denied"))
            as scrape_mock,
        ):
            self.assertEqual(bot.main(), 1)

        scrape_mock.assert_called_once_with(phrases[0])
        self.assertEqual(list(self.output_dir.glob("*.xlsx")), [])

    def test_block_exports_partial_rows_without_overwriting_completed_workbook(self):
        phrases = EXPECTED_SEARCH_TEXTS[:3]
        completed = [{"companyId": "completed", "url": "https://completed.example/"}]
        partial = [
            {"companyId": "partial", "url": "https://partial.example/", "name": "First"},
            {"companyId": "partial", "url": "https://partial.example/", "name": "Duplicate"},
        ]
        self.output_dir.mkdir(parents=True)
        normal_blocked_path = Path(bot.build_output_path(2, phrases[1]))
        bot.pd.DataFrame(completed, columns=EXPECTED_COLUMNS).to_excel(
            normal_blocked_path, index=False, engine="openpyxl"
        )
        previous_completed_bytes = normal_blocked_path.read_bytes()
        with (
            mock.patch.object(bot, "SEARCH_TEXTS", phrases),
            mock.patch.object(bot, "scrape", side_effect=[
                completed, bot.AlibabaAccessBlocked("Access denied", rows=partial)
            ]) as scrape_mock,
        ):
            self.assertEqual(bot.main(), 1)

        self.assertEqual(scrape_mock.call_args_list, [mock.call(phrase) for phrase in phrases[:2]])
        completed_path = Path(bot.build_output_path(1, phrases[0]))
        partial_path = normal_blocked_path.with_name(normal_blocked_path.stem + "_partial.xlsx")
        self.assertEqual(read_workbook_rows(completed_path)[1][:2],
                         ("completed", "https://completed.example/"))
        partial_rows = read_workbook_rows(partial_path)
        self.assertEqual(partial_rows[0], tuple(EXPECTED_COLUMNS))
        self.assertEqual(len(partial_rows), 2)
        self.assertEqual(partial_rows[1][:3], ("partial", "https://partial.example/", "First"))
        self.assertEqual(normal_blocked_path.read_bytes(), previous_completed_bytes)
        self.assertFalse(Path(bot.build_output_path(3, phrases[2])).exists())
        self.assertEqual(len(list(self.output_dir.glob("*.xlsx"))), 3)


if __name__ == "__main__":
    unittest.main()
